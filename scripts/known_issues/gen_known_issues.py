#!/usr/bin/env python3
"""
gen_known_issues.py - Build the Known Issues RST grid table for Astra release
notes from the JIRA sheet of TestRail .xlsx test reports.

* IDs already in previous release notes reuse the previous entry unchanged
  (SoC flags are added for any new SoCs).
* New IDs get a generated entry: Module from the Jira Components (renamed via
  the rename map / previous Module names), Summary from the Jira summary with
  [bracket] tags interpreted (SoCs, profiles, boards) and removed.
* SoC columns default to: SL1620 SL1640 SL1680 SL2611 SL2615 SL2619
  (override with --socs).

Run from the doc/ directory. See scripts/known_issues/README.md.

Examples:
  python3 scripts/known_issues/gen_known_issues.py "reports/*.xlsx" \\
      -p release_notes/scarthgap_6.12_v2.*
  python3 scripts/known_issues/gen_known_issues.py "reports/*.xlsx" \\
      -p release_notes/scarthgap_6.12_v2.* \\
      -r scripts/known_issues/component_renames.json \\
      -t scripts/known_issues/tag_config.json \\
      -u release_notes/scarthgap_6.12_v2.6.0.rst
"""
import argparse
import difflib
import glob
import json
import re
import sys
import textwrap
from dataclasses import dataclass
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")

# --------------------------------------------------------------------------
# Defaults
# --------------------------------------------------------------------------
# Default SoC column order. Override with --socs.
DEFAULT_SOCS = ["SL1620", "SL1640", "SL1680", "SL2611", "SL2615", "SL2619"]

# Jira component -> Module text. Matching ignores case, spaces and punctuation.
# Add/override with -r component_renames.json
DEFAULT_RENAMES = {
    "BT": "Bluetooth",
    "BT,WIFI": "BT and WiFi",
    "WIFI": "WiFi",
    "Gstream Pipeline": "Gstreamer Pipeline",
    "Gstreamer": "Gstreamer Pipeline",
    "Graphics": "Display",
    "Weston_OOBE": "OOBE",
    "Kernel": "Linux Kernel",
    "UBoot": "Bootloader",
    "U-Boot": "Bootloader",
    "PM": "Power Management",
    "Power": "Power Management",
}

# How [bracket] tags in Jira summaries are handled. Extend with -t tag_config.json
DEFAULT_TAG_CONFIG = {
    # Tag SoC -> SoC columns it covers ("x" wildcards like SL16x0 work automatically)
    "soc_aliases": {"SL2610": ["SL2611", "SL2615", "SL2619"]},
    # Profile tag -> note appended to NEW summaries ("" = recognised, no note)
    "profiles": {
        "nand": "NAND profile",
        "spi": "SPI profile",
        "spi_boot": "SPI boot profile",
        "usb_boot": "USB boot profile",
        "oobe": "",
    },
    # Board tag -> note appended to NEW summaries ("" = standard board, no note)
    # RDK = reference design kit (the standard hardware), so no note.
    "boards": {"rdk": ""},
    # Dropped silently (FB = Firebird, internal project name)
    "ignore": ["fb", "firebird"],
    "yocto_releases": ["dunfell", "kirkstone", "langdale", "mickledore", "nanbield",
                       "scarthgap", "styhead", "walnascar", "whinlatter"],
}

ID_RE = re.compile(r"VSSDK-(\d+)", re.I)
SOC_RE = re.compile(r"SL\d{4}", re.I)               # SoC in a report filename
SOC_COL_RE = re.compile(r"^SL\d{4}$", re.I)         # per-SoC column header
SOC_TOKEN_RE = re.compile(r"^sl[\dx]{4}$", re.I)    # SoC inside a tag (SL16x0 ok)
LIST_SOC_HEADERS = {"soc", "socs", "platform", "platforms"}
TAG_RE = re.compile(r"\[([^\]]*)\]")
SEP_RE = re.compile(r"^\+[-=+]+\+$")                # grid-table separator line


def log(*a):
    print(*a, file=sys.stderr)


@dataclass
class Entry:
    ids: list
    socs: set
    module: str
    summary: str
    source: str = ""
    order: int = 0
    report_order: int = 0


def cell(v):
    s = "" if v is None else str(v).strip()
    return "" if s.lower() == "nan" else s


def norm(s):
    return re.sub(r"[^a-z0-9]", "", s.lower())


def tkey(s):
    return re.sub(r"[\s\-]+", "_", s.strip().lower())


def load_json(path):
    data = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    return {k: v for k, v in data.items() if not k.startswith("_")}


def socs_in_text(text, universe):
    """'SL1620, SL16x0' -> {'SL1620', 'SL1640', 'SL1680'}"""
    out = set()
    for tok in re.findall(r"SL[\dX]{4}", text.upper()):
        if "X" in tok:
            pat = re.compile("^" + tok.replace("X", r"\d") + "$")
            out |= {s for s in universe if pat.match(s)}
        else:
            out.add(tok)
    return out


# --------------------------------------------------------------------------
# JIRA sheet
# --------------------------------------------------------------------------
def read_jira_sheet(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    name = next((s for s in wb.sheetnames if s.strip().lower() == "jira"), None)
    if name is None:
        raise SystemExit(f"{path}: no 'JIRA' sheet")
    header, issues = None, []
    for row in wb[name].iter_rows(values_only=True):
        vals = [cell(v) for v in row]
        if header is None:                       # skip title rows ("Known Issue")
            low = [v.lower() for v in vals]
            if "id" in low and "summary" in low:
                header = {k: low.index(k) for k in
                          ("type", "id", "summary", "components", "severity") if k in low}
            continue

        def get(k):
            i = header.get(k)
            return vals[i] if i is not None and i < len(vals) else ""

        m = ID_RE.search(get("id"))
        if m:
            issues.append({"id": m.group(1), "summary": get("summary"),
                           "components": get("components"),
                           "severity": get("severity"), "type": get("type")})
    wb.close()
    if header is None:
        raise SystemExit(f"{path}: JIRA sheet has no ID/Summary header row")
    return issues


# --------------------------------------------------------------------------
# Tags in Jira summaries
# --------------------------------------------------------------------------
def load_tag_config(path):
    cfg = json.loads(json.dumps(DEFAULT_TAG_CONFIG))     # deep copy
    if path:
        for k, v in load_json(path).items():
            if isinstance(v, dict):
                cfg.setdefault(k, {}).update(v)
            elif isinstance(v, list):
                cfg[k] = list(dict.fromkeys(cfg.get(k, []) + v))
            else:
                cfg[k] = v
    return cfg


class TagParser:
    def __init__(self, cfg, universe):
        self.universe = list(universe)
        self.aliases = {k.upper(): [s.upper() for s in v]
                        for k, v in cfg.get("soc_aliases", {}).items()}
        self.profiles = {tkey(k): v for k, v in cfg.get("profiles", {}).items()}
        self.boards = {tkey(k): v for k, v in cfg.get("boards", {}).items()}
        self.ignore = {tkey(k) for k in
                       cfg.get("ignore", []) + cfg.get("yocto_releases", [])}

    def expand_soc(self, tok):
        t = tok.upper()
        if t in self.aliases:
            return set(self.aliases[t])
        if "X" in t:
            pat = re.compile("^" + t.replace("X", r"\d") + "$")
            return {s for s in self.universe if pat.match(s)}
        return {t}

    def parse(self, summary):
        res = {"socs": set(), "notes": [], "profiles": [], "unknown": [],
               "unknown_socs": [], "text": TAG_RE.sub(" ", summary)}
        for tag in TAG_RE.findall(summary):
            toks = [t for t in re.split(r"[\s_\-/,+]+", tag.lower()) if t]
            i = 0
            while i < len(toks):
                pair = f"{toks[i]}_{toks[i + 1]}" if i + 1 < len(toks) else None
                if pair and (pair in self.profiles or pair in self.boards):
                    key, i = pair, i + 2
                else:
                    key, i = toks[i], i + 1
                if SOC_TOKEN_RE.match(key):
                    known = {s for s in self.expand_soc(key) if s in self.universe}
                    res["socs"] |= known
                    if not known:
                        res["unknown_socs"].append(key.upper())
                elif key in self.ignore:
                    continue
                elif key in self.profiles or key in self.boards:
                    if key in self.profiles:
                        res["profiles"].append(key)
                    note = self.profiles.get(key, self.boards.get(key, ""))
                    if note and note not in res["notes"]:
                        res["notes"].append(note)
                else:
                    res["unknown"].append(key)
        return res


def clean_summary(text, notes=()):
    t = text.replace("|", "/")
    t = re.sub(r"\s+", " ", t).strip(" -:;,")
    t = re.sub(r"\s+([.,;:!?])", r"\1", t)             # "cases ." -> "cases."
    if not t:
        t = "TODO: add summary"
    t = t.replace("*", r"\*").replace("@", r"\@")       # RST inline markup
    t = re.sub(r"(\w)_(?=\s|$|[.,;:!?])", r"\1\\_", t)  # accidental RST refs
    t = t[0].upper() + t[1:]
    if t[-1] not in ".!?":
        t += "."
    if notes:
        t += " (" + ", ".join(notes) + ")"
    return t


# --------------------------------------------------------------------------
# Component renaming
# --------------------------------------------------------------------------
class Renamer:
    """Rename map -> exact match to a previous Module -> close match -> as-is."""

    def __init__(self, renames, known_modules, cutoff=0.85, min_len=5):
        self.renames = {norm(k): v for k, v in renames.items()}
        self.known = {}
        for m in list(known_modules) + list(renames.values()):
            if m:
                self.known.setdefault(norm(m), m)
        self.cutoff, self.min_len = cutoff, min_len
        self.log = {}                                   # raw -> (new, how)

    def _note(self, raw, new, how):
        if raw != new or how == "unmapped":
            self.log[raw] = (new, how)
        return new

    def one(self, raw):
        k = norm(raw)
        if not k:
            return ""
        if k in self.renames:
            return self._note(raw, self.renames[k], "rename map")
        if k in self.known:
            return self._note(raw, self.known[k], "previous notes")
        if self.cutoff and len(k) >= self.min_len:
            hit = difflib.get_close_matches(k, list(self.known), n=1, cutoff=self.cutoff)
            if hit:
                return self._note(raw, self.known[hit[0]], "close match - check")
        return self._note(raw, raw, "unmapped")

    def module(self, components):
        if norm(components) in self.renames:            # whole value, e.g. "BT,WIFI"
            return self._note(components, self.renames[norm(components)], "rename map")
        names = []
        for part in re.split(r"[,;/]", components):
            n = self.one(part.strip())
            if n and n not in names:
                names.append(n)
        return " and ".join(names) or "TBD"


# --------------------------------------------------------------------------
# Previous release notes (RST grid tables)
# --------------------------------------------------------------------------
def is_row(line):
    s = line.strip()
    return bool(s) and (s[0] in "+|" or "|" in s)


def find_tables(lines):
    tables, i, n = [], 0, len(lines)
    while i < n:
        if SEP_RE.match(lines[i].strip()):
            j = i + 1
            while j < n and is_row(lines[j]):
                j += 1
            tables.append((i, j))
            i = j
        else:
            i += 1
    return tables


def split_cells(line, ncols):
    # Split on '|' (not by position) so misaligned hand-edited rows, or rows
    # missing the leading '|', still parse.
    s = line.strip()
    s = s[1:] if s.startswith("|") else s
    s = s[:-1] if s.endswith("|") else s
    parts = s.split("|")
    if len(parts) > ncols:
        parts = parts[:ncols - 1] + ["|".join(parts[ncols - 1:])]
    parts += [""] * (ncols - len(parts))
    return [p.strip() for p in parts]


_order = [0]


def parse_known_issues_table(tlines, source="", universe=DEFAULT_SOCS):
    """Return (soc_columns, [Entry], ignored_headers) or None."""
    ncols = tlines[0].strip().count("+") - 1
    if ncols < 3:
        return None
    groups, cur, header_end = [], [], None
    for ln in tlines[1:]:
        s = ln.strip()
        if SEP_RE.match(s):
            groups.append(cur)
            cur = []
            if "=" in s and header_end is None:
                header_end = len(groups)
        else:
            cur.append(split_cells(ln, ncols))
    if cur:
        groups.append(cur)
    if header_end is None:
        return None

    def merge(rows):
        return [[r[c] for r in rows if r[c]] for c in range(ncols)]

    header = [" ".join(c).strip()
              for c in merge([r for g in groups[:header_end] for r in g])]
    low = [h.lower() for h in header]
    if not {"module", "id", "summary"} <= set(low):
        return None
    i_mod, i_id, i_sum = low.index("module"), low.index("id"), low.index("summary")

    # Only headers that look like a SoC (SL1620) are SoC flag columns.
    # A "SoC"/"Platform" column is read as a list of SoCs. Others are ignored.
    soc_cols, list_cols, ignored = {}, [], []
    for k, h in enumerate(header[:i_mod]):
        if SOC_COL_RE.match(h):
            soc_cols[k] = h.upper()
        elif h.lower() in LIST_SOC_HEADERS:
            list_cols.append(k)
        elif h:
            ignored.append(h)

    entries = []
    for g in groups[header_end:]:
        if not g:
            continue
        cols = merge(g)
        ids = re.findall(r"\d+", " ".join(cols[i_id]))
        if not ids:
            continue
        socs = {name for k, name in soc_cols.items()
                if "Y" in " ".join(cols[k]).upper().split()}
        for k in list_cols:
            socs |= socs_in_text(" ".join(cols[k]), universe)
        _order[0] += 1
        entries.append(Entry(ids, socs, " ".join(cols[i_mod]),
                             " ".join(cols[i_sum]), source, _order[0]))
    return list(soc_cols.values()), entries, ignored


def load_previous(paths, universe):
    """Later files win for the same ID."""
    db, warnings, files = {}, [], []
    for p in map(Path, paths):
        files += sorted(p.rglob("*.rst")) if p.is_dir() else [p]
    for p in files:
        if not p.exists():
            warnings.append(f"previous file not found: {p}")
            continue
        lines = p.read_text(encoding="utf-8").splitlines()
        for s, e in find_tables(lines):
            parsed = parse_known_issues_table(lines[s:e], str(p), universe)
            if not parsed:
                continue
            _cols, entries, ignored = parsed
            for h in ignored:
                warnings.append(f"{p}: ignored column '{h}' (not a SoC column)")
            for ent in entries:
                for jid in ent.ids:
                    db[jid] = ent
    return db, warnings


# --------------------------------------------------------------------------
# Build + render
# --------------------------------------------------------------------------
def build(reports, db, renamer, tags, fresh_socs=False, use_tag_socs=True):
    entries, new, seq = {}, [], 10 ** 9
    seen = 0                      # order in which the reports mention an issue
    flags_added = []
    rep = {"unknown": {}, "unknown_socs": {}, "mismatch": [], "missing_profile": []}

    for _path, soc, issues in reports:
        for iss in issues:
            jid = iss["id"]
            t = tags.parse(iss["summary"])
            for u in t["unknown"]:
                rep["unknown"].setdefault(u, set()).add(jid)
            for u in t["unknown_socs"]:
                rep["unknown_socs"].setdefault(u, set()).add(jid)
            if t["socs"] and soc not in t["socs"]:
                rep["mismatch"].append((jid, soc, sorted(t["socs"])))
            add_socs = {soc} | (t["socs"] if use_tag_socs else set())

            prev = db.get(jid)
            if prev:
                key = tuple(prev.ids)
                if key not in entries:
                    seen += 1
                    entries[key] = Entry(list(prev.ids),
                                         set() if fresh_socs else set(prev.socs),
                                         prev.module, prev.summary,
                                         prev.source, prev.order, seen)
                ent = entries[key]
                for s in sorted(add_socs):
                    if s not in ent.socs:
                        ent.socs.add(s)
                        if not fresh_socs and (jid, s) not in flags_added:
                            flags_added.append((jid, s))
                for pk in t["profiles"]:
                    note = tags.profiles.get(pk)
                    word = pk.split("_")[0]
                    if note and word not in ent.summary.lower() and \
                            (jid, note) not in rep["missing_profile"]:
                        rep["missing_profile"].append((jid, note))
            else:
                key = (jid,)
                if key not in entries:
                    seq += 1
                    seen += 1
                    entries[key] = Entry([jid], set(),
                                         renamer.module(iss["components"]),
                                         clean_summary(t["text"], t["notes"]),
                                         "new", seq, seen)
                    new.append((entries[key], iss))
                entries[key].socs |= add_socs
    return list(entries.values()), new, flags_added, rep


def render(entries, soc_cols, width=83):
    headers = soc_cols + ["Module", "ID", "Summary"]
    rows = []
    for e in entries:
        ids = []
        for k, jid in enumerate(e.ids):
            ids += ([""] if k else []) + [jid]
        summ = textwrap.wrap(e.summary, width, break_long_words=False,
                             break_on_hyphens=False) or [""]
        rows.append([["Y" if s in e.socs else "N/A"] for s in soc_cols]
                    + [[e.module], ids, summ])
    widths = [max([len(h)] + [len(l) for r in rows for l in r[i]])
              for i, h in enumerate(headers)]
    nsoc = len(soc_cols)

    def sep(ch):
        return "+" + "+".join(ch * (w + 2) for w in widths) + "+"

    def line(cells):
        return "|" + "|".join(
            " " + (t.center(w) if i < nsoc else t.ljust(w)) + " "
            for i, (t, w) in enumerate(zip(cells, widths))) + "|"

    out = [sep("-"), line(headers), sep("=")]
    for r in rows:
        for k in range(max(len(c) for c in r)):
            out.append(line([c[k] if k < len(c) else "" for c in r]))
        out.append(sep("-"))
    return "\n".join(out)


def replace_table(path, table, universe):
    p = Path(path)
    lines = p.read_text(encoding="utf-8").splitlines()
    for s, e in find_tables(lines):
        if parse_known_issues_table(lines[s:e], universe=universe):
            indent = re.match(r"\s*", lines[s]).group(0)
            lines[s:e] = [indent + l for l in table.splitlines()]
            p.write_text("\n".join(lines) + "\n", encoding="utf-8")
            return True
    return False


# --------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("reports", nargs="+", help="test report .xlsx files (globs ok)")
    ap.add_argument("-p", "--previous", nargs="*", default=[],
                    help="previous release-notes .rst files or directories")
    ap.add_argument("-u", "--update",
                    help="replace the known-issues table in this .rst in place "
                         "(its current table is also used as previous notes)")
    ap.add_argument("-o", "--output", help="write the table to this file")
    ap.add_argument("--soc", help="SoC to use if it isn't in the report filename")
    ap.add_argument("--socs",
                    help="comma-separated SoC columns, in order "
                         f"(default: {','.join(DEFAULT_SOCS)})")
    ap.add_argument("-r", "--renames", "--component-map", dest="renames",
                    help='JSON file of component renames: {"Gstream Pipeline": "Gstreamer Pipeline"}')
    ap.add_argument("-t", "--tags", help="JSON tag config (soc_aliases/profiles/boards/ignore)")
    ap.add_argument("--match-cutoff", type=float, default=0.85,
                    help="similarity (0-1) for automatic close matches; 0 disables")
    ap.add_argument("--no-tag-socs", action="store_true",
                    help="don't set SoC flags from [SLxxxx] tags in Jira summaries")
    ap.add_argument("--fresh-socs", action="store_true",
                    help="set SoC flags only from the current reports")
    ap.add_argument("--sort", choices=["report", "previous", "id"], default="report",
                    help="report: SoC column order, then report row order; "
                         "previous: keep the previous notes' order; id: by Jira ID")
    ap.add_argument("--width", type=int, default=83, help="Summary column width")
    args = ap.parse_args()

    # SoC columns: fixed default order unless overridden
    if args.socs:
        soc_cols = [s.strip().upper() for s in args.socs.split(",") if s.strip()]
    else:
        soc_cols = list(DEFAULT_SOCS)
    universe = list(dict.fromkeys(DEFAULT_SOCS + soc_cols))

    # Reports
    paths = []
    for r in args.reports:                       # expand globs (Windows cmd)
        paths += sorted(glob.glob(r)) if any(c in r for c in "*?[") else [r]
    if not paths:
        sys.exit("No report files found")
    reports = []
    for r in paths:
        m = SOC_RE.search(Path(r).name)
        soc = (args.soc or (m.group(0) if m else "")).upper()
        if not soc:
            sys.exit(f"{r}: no SoC in filename; pass --soc")
        if soc not in universe:
            universe.append(soc)
        issues = read_jira_sheet(r)
        reports.append((r, soc, issues))
        log(f"{soc}: {len(issues)} Jira issue(s) in {Path(r).name}")

    # Process the reports in SoC column order (SL1620, then SL1640, ...) so new
    # entries are added to the table in that order.
    def soc_rank(item):
        soc = item[1]
        order = soc_cols if soc in soc_cols else universe
        rank = order.index(soc) if soc in order else len(order)
        return (0 if soc in soc_cols else 1, rank, Path(item[0]).name)

    reports.sort(key=soc_rank)
    log(f"Report order: {' | '.join(s for _p, s, _i in reports)}")

    # Previous notes (the --update file is read last, so hand edits win)
    prev_paths = list(args.previous) + ([args.update] if args.update else [])
    db, prev_warnings = load_previous(prev_paths, universe)
    log(f"Previous notes: {len({tuple(e.ids) for e in db.values()})} entries loaded")

    renames = dict(DEFAULT_RENAMES)
    if args.renames:
        renames.update(load_json(args.renames))
    renamer = Renamer(renames, {e.module for e in db.values()}, args.match_cutoff)
    tags = TagParser(load_tag_config(args.tags), universe)

    entries, new, flags_added, rep = build(reports, db, renamer, tags,
                                           args.fresh_socs, not args.no_tag_socs)

    # SoCs outside the column list
    extras = sorted({s for e in entries for s in e.socs} - set(soc_cols))
    dropped = []
    if extras:
        if args.socs:
            keep = []
            for e in entries:
                (keep if e.socs & set(soc_cols) else dropped).append(e)
            entries = keep
        else:
            soc_cols = soc_cols + extras

    sort_key = {"id": lambda e: int(e.ids[0]),
                "previous": lambda e: e.order,
                "report": lambda e: e.report_order}[args.sort]
    entries.sort(key=sort_key)
    table = render(entries, soc_cols, args.width)

    # ---------------- review report ----------------
    log("")
    log(f"SoC columns: {' | '.join(soc_cols)}")
    log(f"Reused from previous notes: {len(entries) - len([n for n in new if n[0] in entries])}")
    log(f"New entries (review module/summary): {len(new)}")
    for e, iss in new:
        log(f"  {e.ids[0]}  [{e.module}]  {e.summary}  "
            f"(Components: {iss['components'] or '-'}, Severity: {iss['severity'] or '-'})")
    if flags_added:
        log("SoC flags added to reused entries:")
        for jid, soc in flags_added:
            log(f"  {jid} -> {soc}")
    if renamer.log:
        log("Component renames:")
        for raw, (name, how) in sorted(renamer.log.items()):
            log(f"  {raw!r:28} -> {name!r:28} ({how})")
    if rep["unknown"]:
        log("Unrecognised summary tags (add to 'profiles', 'boards' or 'ignore' in the tag file):")
        for tag, ids in sorted(rep["unknown"].items()):
            log(f"  [{tag}]  {', '.join(sorted(ids))}")
    if rep["unknown_socs"]:
        log("SoC tags that don't match a column (add to 'soc_aliases'):")
        for tag, ids in sorted(rep["unknown_socs"].items()):
            log(f"  [{tag}]  {', '.join(sorted(ids))}")
    if rep["mismatch"]:
        log("SoC mismatch (issue in a report whose SoC isn't in its tags):")
        for jid, soc, tagged in rep["mismatch"]:
            log(f"  {jid}: report {soc}, tags {', '.join(tagged)}")
    if rep["missing_profile"]:
        log("Reused entries whose summary doesn't mention a tagged profile:")
        for jid, note in rep["missing_profile"]:
            log(f"  {jid}: ({note})")
    if extras and not args.socs:
        log(f"Extra SoC column(s) added (not in default list): {', '.join(extras)}")
    if extras and args.socs:
        log(f"SoC(s) not in --socs, flags omitted: {', '.join(extras)}")
    if dropped:
        log("Rows omitted (no SoC in --socs columns):")
        for e in dropped:
            log(f"  {'/'.join(e.ids)}  {e.summary}")
    for w in prev_warnings:
        log(f"warning: {w}")

    # ---------------- output ----------------
    if args.update:
        if not replace_table(args.update, table, universe):
            sys.exit(f"No known-issues table found in {args.update}")
        log(f"Updated {args.update}")
    if args.output:
        Path(args.output).write_text(table + "\n", encoding="utf-8")
        log(f"Wrote {args.output}")
    if not args.update and not args.output:
        print(table)


if __name__ == "__main__":
    main()
